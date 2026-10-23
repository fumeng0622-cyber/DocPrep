
import unittest
from io import BytesIO

from docx import Document

from word_parser import parse_docx


def create_test_docx() -> bytes:
    """在内存中创建测试 Word 文件，不需要额外准备样本。"""

    document = Document()

    document.add_paragraph("公文格式要求")
    document.add_paragraph("正文一般使用仿宋，三号字。")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "规范名称"
    table.cell(0, 1).text = "字号"
    table.cell(1, 0).text = "正文"
    table.cell(1, 1).text = "三号"

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


class WordParserTests(unittest.TestCase):

    def setUp(self):
        self.file_bytes = create_test_docx()

    def test_extracts_paragraphs(self):
        chunks = parse_docx(self.file_bytes, "测试.docx")

        contents = [chunk.content for chunk in chunks]

        self.assertIn("公文格式要求", contents)
        self.assertIn("正文一般使用仿宋，三号字。", contents)

    def test_extracts_table_rows(self):
        chunks = parse_docx(self.file_bytes, "测试.docx")

        table_chunks = [
            chunk for chunk in chunks
            if chunk.metadata.get("element_type") == "table_row"
        ]

        self.assertEqual(len(table_chunks), 2)
        self.assertEqual(
            table_chunks[0].content,
            "规范名称 | 字号",
        )
        self.assertEqual(
            table_chunks[1].content,
            "正文 | 三号",
        )

    def test_preserves_source_information(self):
        chunks = parse_docx(self.file_bytes, "测试.docx")

        self.assertTrue(chunks)

        for chunk in chunks:
            self.assertEqual(chunk.source_file, "测试.docx")
            self.assertEqual(chunk.source_type, "docx")

    def test_empty_document_returns_empty_list(self):
        document = Document()
        buffer = BytesIO()
        document.save(buffer)

        chunks = parse_docx(buffer.getvalue(), "空文档.docx")

        self.assertEqual(chunks, [])


if __name__ == "__main__":
    unittest.main()
