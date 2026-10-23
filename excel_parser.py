
from io import BytesIO

from openpyxl import load_workbook

from knowledge_schema import KnowledgeChunk


def parse_xlsx(file_bytes, filename):
    """
    解析 Excel 文件，将数据行转换为知识块。

    file_bytes: Excel 文件的二进制内容
    filename: 原始文件名

    返回：KnowledgeChunk 列表
    """

    workbook = load_workbook(
        filename=BytesIO(file_bytes),
        read_only=True,
        data_only=True
    )

    chunks = []

    try:
        for sheet in workbook.worksheets:
            # 获取工作表中的所有行
            rows = sheet.iter_rows(values_only=True)

            # 第一行作为表头
            header_row = next(rows, None)

            if header_row is None:
                continue

            headers = []

            for index, value in enumerate(header_row):
                if value is None or str(value).strip() == "":
                    headers.append(f"列{index + 1}")
                else:
                    headers.append(str(value).strip())

            # 从第二行开始读取数据
            for row_number, row in enumerate(rows, start=2):

                # 跳过完全为空的行
                if all(value is None for value in row):
                    continue

                content_lines = []

                for header, value in zip(headers, row):
                    if value is None:
                        continue

                    content_lines.append(
                        f"{header}：{value}"
                    )

                if not content_lines:
                    continue

                content = "\n".join(content_lines)

                chunks.append(
                    KnowledgeChunk(
                        content=content,
                        source_file=filename,
                        source_type="xlsx",
                        title=sheet.title,
                        section=sheet.title,
                        metadata={
                            "element_type": "excel_row",
                            "sheet_name": sheet.title,
                            "row_number": row_number,
                            "column_count": len(headers),
                        }
                    )
                )

    finally:
        workbook.close()

    return chunks
